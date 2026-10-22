"""build_report.py -- builds the assignment report (7-8 pages) from the real runs."""

import pathlib

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = pathlib.Path(__file__).resolve().parent
SHOTS = ROOT / "shots"
OUT = ROOT / "Report_InterProcessCommunication.docx"

INK = RGBColor(0x00, 0x00, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)
CODE_FILL = "F5F5F5"

STUDENT = "Mohammad Istahad Alam Chowdhury"
NIM = "20250040225"
CLASS = "TI25I"
PROGRAM = "Informatics Engineering"
UNIVERSITY = "Universitas Nusa Putra"
COURSE = "Parallel and Distributed Systems"
LECTURER = "Muchtar Ali Setyo Yudono, S.T., M.T."
DATE = "22 October 2026"


# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------
def code_block(doc, text, size=8.0):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(4)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.0
    pf.left_indent = Cm(0.3)

    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "4")
        el.set(qn("w:color"), "DDDDDD")
        borders.append(el)
    pPr.append(borders)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), CODE_FILL)
    pPr.append(shd)

    lines = text.strip("\n").split("\n")
    for i, line in enumerate(lines):
        run = p.add_run(line if line.strip() else " ")
        run.font.name = "Courier New"
        run.font.size = Pt(size)
        if i < len(lines) - 1:
            run.add_break()
    return p


def table(doc, headers, rows, widths, size=9.0):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, head in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(head)
        run.bold = True
        run.font.size = Pt(size)
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(str(value)).font.size = Pt(size)
    for r in t.rows:
        for i, width in enumerate(widths):
            r.cells[i].width = Cm(width)
            # no single row may be split across pages
            trPr = r._tr.get_or_add_trPr()
            cant = OxmlElement("w:cantSplit")
            cant.set(qn("w:val"), "true")
            trPr.append(cant)

    # keep the whole table on one page: every paragraph but the last row asks
    # to stay with the next one
    for row in t.rows[:-1]:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(doc, path, caption, width_cm=12.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True      # never leave the caption alone
    p.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(9)
    run = cap.add_run(caption)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = GREY


def para(doc, text, size=10.5, space_after=6, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.add_run(text).font.size = Pt(size)
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    return h


def centered(doc, text, size, bold=False, space_after=2, colour=INK, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    return p


def hrule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(10)
    pPr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "333333")
    borders.append(bottom)
    pPr.append(borders)


def field(run, instruction):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
def build():
    doc = Document()

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.different_first_page_header_footer = True

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    for level, size in ((1, 13), (2, 11)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = INK
        style.font.bold = True
        style.font.italic = False
        style.paragraph_format.space_before = Pt(12 if level == 1 else 9)
        style.paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = GREY
    field(run, "PAGE")

    # ---------------- cover ----------------
    for _ in range(4):
        doc.add_paragraph()

    centered(doc, "ASSIGNMENT REPORT", 11, space_after=4, colour=GREY)
    centered(doc, "Inter-Process Communication", 24, bold=True, space_after=2)
    centered(doc, "A client and a server program in Python using sockets", 12,
             colour=GREY, space_after=10)
    hrule(doc)

    centered(doc, COURSE, 12, bold=True, space_after=1)
    centered(doc, "Session 3", 11, colour=GREY, space_after=18)

    centered(doc, "Written by", 10, colour=GREY, space_after=1)
    centered(doc, STUDENT, 14, bold=True, space_after=1)
    centered(doc, f"NIM {NIM}", 11, space_after=1)
    centered(doc, f"Class {CLASS}", 11, space_after=14)

    centered(doc, PROGRAM, 11, space_after=1)
    centered(doc, UNIVERSITY, 11, bold=True, space_after=18)

    centered(doc, f"Lecturer: {LECTURER}", 10.5, space_after=1)
    centered(doc, DATE, 10.5, colour=GREY)

    page_break(doc)

    # ---------------- 1. Introduction ----------------
    heading(doc, "1. Introduction", 1)
    para(doc, "A distributed system is made of programs that run separately and "
              "still need to exchange data. Before any two machines can work "
              "together they need a channel to send bytes through, and sockets "
              "are the usual way to build that channel in Python. A socket is "
              "just an address (IP address plus port number) that one program "
              "can listen on and another program can connect to.")
    para(doc, "This assignment asks for two programs. The server listens on a "
              "port, waits for a client, prints whatever the client sends, and "
              "answers with a confirmation before closing the connection and "
              "waiting for the next client. The client connects to that server "
              "and lets the user type messages one after another, printing the "
              "reply each time. Both programs only use the standard library, so "
              "there was nothing to install.")
    para(doc, "I wrote and tested everything on Windows with Python 3.13.5. All "
              "the screenshots and numbers in this report come from real runs of "
              "the two programs.")

    # ---------------- 2. How it works ----------------
    heading(doc, "2. How the two programs talk to each other", 1)
    para(doc, "The server creates a socket, binds it to port 5000, and calls "
              "listen(). From then on the operating system queues any incoming "
              "connection until the program calls accept(). When a client calls "
              "connect(), the operating system completes a TCP handshake and "
              "accept() returns a new socket that belongs only to that client. "
              "The original listening socket stays open, which is why the server "
              "can carry on serving clients after the first one leaves.")
    figure(doc, SHOTS / "fig_arch.png",
           "Figure 1. The client and the server are separate processes joined by "
           "one TCP connection on port 5000.", width_cm=12.0)
    para(doc, "TCP gives a stream of bytes, not a stream of messages. If the "
              "client sends three messages quickly, the server may receive the "
              "first in two pieces or the first two messages in one read. To keep "
              "messages apart both programs end every message with a newline "
              "character and keep reading until that newline arrives. Figure 2 "
              "shows one full exchange between the two programs.")
    figure(doc, SHOTS / "fig_sequence.png",
           "Figure 2. The calls made during one message and its confirmation.",
           width_cm=12.0)

    # ---------------- 3. Server ----------------
    heading(doc, "3. The server program (Task 1)", 1)
    para(doc, "The main loop of server.py accepts one client at a time, which is "
              "what the assignment asks for. Each connection gets its own socket, "
              "so closing it after the reply does not disturb the listening "
              "socket. The reply also tells the client which message it was.")
    code_block(doc,
               'server = socket.socket(socket.AF_INET, socket.SOCK_STREAM)\n'
               'server.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)\n'
               'server.bind((host, port))\n'
               'server.listen(BACKLOG)\n'
               '\n'
               'while True:\n'
               '    conn, addr = server.accept()\n'
               '    with conn:\n'
               '        handle_client(conn, addr, client_no, timeout, max_messages)')
    para(doc, "For every message the server prints the text, the sender's "
              "address, the length in characters and bytes, and the time it "
              "arrived. It then sends back \"Message received (43 chars / 43 "
              "bytes) - thank you, client #1\" and closes the connection. A "
              "timeout of 30 seconds makes sure a client that connects and then "
              "stays silent cannot block everyone behind it.")
    figure(doc, SHOTS / "fig2_server.png",
           "Figure 3. The server terminal after three clients were served one "
           "after another, each with its own number.", width_cm=12.5)

    # ---------------- 4. Client ----------------
    heading(doc, "4. The client program (Task 2)", 1)
    para(doc, "The client connects once and then loops on input(), so the user "
              "can keep sending messages without restarting the program. For each "
              "line it sends the text, waits for the answer, and prints the round "
              "trip time measured with time.perf_counter().")
    code_block(doc,
               'while True:\n'
               '    message = input("\\nYou > ").strip()\n'
               '    if not message:\n'
               '        continue\n'
               '    if message.lower() in QUIT_COMMANDS:      # /quit, /exit, :q\n'
               '        break\n'
               '    sock, reply, reconnected = exchange(host, port, sock, message)\n'
               '    print(f"Server > {reply}")')
    para(doc, "One thing needed extra work. The server closes the connection "
              "after every reply, so the second message would be sent to a socket "
              "that no longer exists. The client therefore checks whether the "
              "reply came back cleanly, and if the connection is gone it opens a "
              "new one and sends the message again. The user does not notice "
              "anything except a short note in the output.")
    figure(doc, SHOTS / "fig1_client.png",
           "Figure 4. Three messages typed in the client, with the reply and the "
           "round trip time for each.", width_cm=14.0)

    # ---------------- 5. How to run ----------------
    heading(doc, "5. How to run the two programs", 1)
    para(doc, "Both files sit in the same folder. The server has to be started "
              "first, in its own terminal window; the client then goes in a second "
              "window. If the client is started first it prints a message saying "
              "the connection was refused.")
    code_block(doc,
               'python server.py              (first window)\n'
               'python client.py              (second window, same port)')
    para(doc, "On Windows the folder also contains start_demo.bat, which opens "
              "both windows with the correct port already filled in. In the "
              "client window every line typed and confirmed with Enter is sent to "
              "the server, and the answer appears underneath it. Typing /quit "
              "closes the connection and the server keeps waiting for the next "
              "client until Ctrl+C is pressed.")
    para(doc, "The server listens on port 5000 and accepts other machines on the "
              "same network, so the client can be pointed at the server's address "
              "with python client.py --host 192.168.1.10. It is worth saying that "
              "opening the port in a web browser does not work and is not supposed "
              "to: the browser would send an HTTP request, while this server only "
              "understands the newline protocol from section 2. The client program "
              "is the only thing that can talk to it.")

    # ---------------- 6. Testing ----------------
    heading(doc, "6. Testing", 1)
    para(doc, "I tested with a script that starts the real server on a free port, "
              "sends real messages over a real socket, and compares the answers. "
              "That way the tests do not depend on me typing correctly. All seven "
              "cases pass.")
    table(
        doc,
        ["#", "What was tested", "Expected", "Result"],
        [
            ["1", "One client sends one message", "Reply starts with \"Message received\"", "Pass"],
            ["2", "Five clients one after another", "Each served, numbered in order", "Pass"],
            ["3", "Text with accents and an emoji", "45 characters counted correctly", "Pass"],
            ["4", "A message of 5000 characters", "Received in full, not cut short", "Pass"],
            ["5", "A client that connects and leaves", "Server keeps working", "Pass"],
            ["6", "An empty message", "No reply, connection closed", "Pass"],
            ["7", "Four clients at the same time (bonus server)", "All four answered in 0.50 s", "Pass"],
        ],
        widths=[0.9, 6.2, 6.4, 1.8],
    )
    figure(doc, SHOTS / "fig3_tests.png",
           "Figure 5. The test script reporting 7 of 7 cases passed.", width_cm=13.5)

    # ---------------- 7. Problems ----------------
    heading(doc, "7. Problems I ran into", 1)
    para(doc, "Three problems were not obvious at first and each one only showed "
              "up because I ran the programs instead of only reading the code.")
    para(doc, "The first was WinError 10053. My first client worked once and then "
              "crashed on the second message, because the server had already "
              "closed the connection. The fix is the reconnect logic in section 4.")
    para(doc, "The second was an encoding error. The server stopped completely "
              "when a message contained an emoji, with UnicodeEncodeError from the "
              "charmap codec. The network part was fine; the problem was that a "
              "Windows console still uses an old code page by default. Setting the "
              "standard output to UTF-8 fixed it, and the same change on the input "
              "side stopped accented letters from arriving garbled.")
    para(doc, "The third was a truncated message. A single recv() call is not "
              "guaranteed to return a whole message, and in testing a long "
              "message arrived in several pieces. Reading in a loop until the "
              "newline turned up fixed it, which is also why test 4 exists.")
    para(doc, "Apart from those, I added two small things while testing: "
              "SO_REUSEADDR, because restarting the server sometimes failed with "
              "\"address already in use\", and sendall() instead of send(), "
              "because send() can accept only part of a long message.")
    para(doc, "I also wrote a second version of the server that handles each "
              "client in its own thread, as an extra. With four clients that each "
              "kept the connection open for half a second, the threaded server "
              "finished in 0.50 seconds. A sequential server cannot overlap them, "
              "so it would need at least two seconds for the same work.")

    # ---------------- 8. Conclusion ----------------
    heading(doc, "8. Conclusion", 1)
    para(doc, "The two programs do what the assignment asked. The server accepts "
              "clients one at a time, prints every message it receives, answers "
              "with a confirmation and closes the connection properly. The client "
              "connects, sends typed messages in a loop and shows the replies with "
              "their round trip times. Seven automated tests pass, including a "
              "5000-character message and text with non-ASCII characters.")
    para(doc, "The part I did not expect was how much of the work is deciding "
              "where a message ends and what to do when the other side hangs up. "
              "TCP only delivers bytes, so the newline rule had to be written by "
              "hand, and the server closing after each reply forced the client to "
              "reconnect. Testing each of those cases was what turned the "
              "assignment from something that looked right into something that "
              "actually works.")

    # ---------------- References ----------------
    heading(doc, "References", 1)
    refs = [
        "Python Software Foundation. \"socket â€” Low-level networking "
        "interface.\" Python 3 documentation. "
        "https://docs.python.org/3/library/socket.html",
        "Kurose, J. F. and Ross, K. W. Computer Networking: A Top-Down "
        "Approach, 8th edition. Pearson, 2021.",
        "Tanenbaum, A. S. and van Steen, M. Distributed Systems, 4th edition, "
        "version 4.02. 2023.",
        "Stevens, W. R., Fenner, B. and Rudoff, A. M. UNIX Network Programming, "
        "Volume 1: The Sockets Networking API, 3rd edition. Addison-Wesley, 2003.",
        "RFC 9293. Transmission Control Protocol (TCP). IETF, 2022.",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.add_run(f"[{i}]  {ref}").font.size = Pt(9.5)

    doc.core_properties.title = "Inter-Process Communication"
    doc.core_properties.author = STUDENT
    doc.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    build()


